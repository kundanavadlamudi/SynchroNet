"""Generate a professional Word document (.docx) containing the comprehensive
NexTrace Engineering Continuation and Completion Report.
"""

from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

ROOT = Path(__file__).resolve().parent

def set_cell_background(cell, hex_color: str):
    """Set background color of a table cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set internal cell padding."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_report():
    doc = Document()

    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Color Palette
    PRIMARY = RGBColor(14, 116, 144)      # Cyan / Deep Teal
    SECONDARY = RGBColor(124, 58, 237)   # Violet
    DARK_TEXT = RGBColor(30, 41, 59)     # Slate 800
    MUTED_TEXT = RGBColor(100, 116, 139) # Slate 500
    ACCENT_RED = RGBColor(220, 38, 38)

    # Base Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = DARK_TEXT

    # -------------------------------------------------------------
    # Cover / Header Section
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_eyebrow = title_p.add_run("SOCIAL MEDIA INTELLIGENCE PLATFORM\n")
    run_eyebrow.font.size = Pt(9.5)
    run_eyebrow.font.bold = True
    run_eyebrow.font.color.rgb = PRIMARY

    run_title = title_p.add_run("NexTrace — Engineering Continuation & System Completion Report")
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = DARK_TEXT

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(4)
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run("Comprehensive Audit, Data Contracts, Ingestion Architecture, Resilient NLP, Pipeline Unification, Model Audit, and Dashboard Integration")
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = MUTED_TEXT

    # Metadata Card Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_data = [
        ("Repository", "https://github.com/kundanavadlamudi/SynchroNet.git (Branch: main)"),
        ("Date & Status", "October 24, 2026 | All 42 Tests Passed (0 Regressions, 0 Failures)"),
        ("Role", "Senior Python / ML Systems Engineer"),
        ("Core Artifacts", "models/selected_bot_classifier.joblib (8.9 MB Preserved Intact)"),
    ]

    for i, (label, val) in enumerate(meta_data):
        row = meta_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(1.8)
        c1.width = Inches(4.7)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "F8FAFC")
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = PRIMARY

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.size = Pt(10)
        r1.font.color.rgb = DARK_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Disclaimer Callout Box
    box_table = doc.add_table(rows=1, cols=1)
    box_cell = box_table.rows[0].cells[0]
    box_cell.width = Inches(6.5)
    set_cell_background(box_cell, "FEF2F2")
    set_cell_margins(box_cell, 120, 120, 150, 150)
    p_box = box_cell.paragraphs[0]
    p_box.paragraph_format.space_after = Pt(0)
    r_warn = p_box.add_run("IMPORTANT ETHICAL & OPERATIONAL DISCLAIMER:\n")
    r_warn.font.bold = True
    r_warn.font.size = Pt(10)
    r_warn.font.color.rgb = ACCENT_RED
    r_warn_text = p_box.add_run(
        "NexTrace generates investigation leads, NOT proof of disinformation, malicious intent, or coordinated manipulation. "
        "A high coordination score or classifier probability reflects statistical resemblance to historical benchmark patterns. "
        "Every flagged account pair and candidate campaign requires human investigator review before any conclusions are drawn."
    )
    r_warn_text.font.size = Pt(9.5)
    r_warn_text.font.color.rgb = DARK_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # -------------------------------------------------------------
    # 1. Executive Summary & Audit Findings
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1_run = h1.add_run("1. Executive Summary & Audit Findings")
    h1_run.font.color.rgb = PRIMARY

    p = doc.add_paragraph(
        "An exhaustive initial audit was conducted against the repository prior to code modifications. "
        "The findings established an empirical baseline, verified prior reported status, and isolated critical blockers:"
    )

    doc.add_paragraph(
        "• Baseline Test Reproduction: The prior-reported result of 27 passing tests was independently verified and passed (pytest 8.4.2 in 25.46s).",
        style='List Bullet'
    )
    doc.add_paragraph(
        "• Model Artifact Preservation: models/selected_bot_classifier.joblib (8.9 MB) was inspected and verified intact. The artifact is a scikit-learn pipeline bundle (SimpleImputer + RandomForestClassifier) trained on 16,318 accounts from Cresci-2015 and Cresci-2017.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "• Critical Pipeline Blocker: run_pipeline.py failed at Stage 4 (Semantic similarity) due to an unprotected top-level import of sentence_transformers. In environments without PyTorch installed, the entire pipeline broke.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "• Ingestion Gap: The system had no file or API ingestion boundary, relying exclusively on 10 hardcoded synthetic posts.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "• Disconnected ML: Account behavioral features extracted in Phase 1 were never routed to the Phase 2 bot classifier.",
        style='List Bullet'
    )

    # -------------------------------------------------------------
    # 2. Reliable Data Contracts (Stage 2)
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1_run = h1.add_run("2. Reliable Unified Data Contracts (Stage 2)")
    h1_run.font.color.rgb = PRIMARY

    doc.add_paragraph(
        "To ensure system stability across varying data sources, src/contracts.py was implemented. "
        "It defines strict typing, date normalization, and validation boundaries so malformed records "
        "are safely filtered without crashing the pipeline:"
    )

    contract_table = doc.add_table(rows=5, cols=3)
    contract_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Contract Entity", "Core Schema Fields", "Handling of Malformed & Missing Data"]
    for j, h in enumerate(headers):
        cell = contract_table.rows[0].cells[j]
        set_cell_background(cell, "0E7490")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(10)

    rows_data = [
        ("PostRecord", "post_id, account_id, text, timestamp, url, source_platform, provenance, collection_time", "Deduplicates post_id; converts timestamps to UTC; strips whitespace; drops blank IDs/text with warning logs."),
        ("AccountRecord", "account_id, followers_count, friends_count, statuses_count, favourites_count, default_profile, geo_enabled", "Numeric features coerced to integer; missing metadata neutrally defaulted to 0 to avoid false bot penalties."),
        ("AnalysisResultRecord", "subject_1, subject_2, post_1, post_2, coordination_score, investigation_priority, evidence, limitations", "Enforces valid float scores (0.0 to 1.0); priority categorized as High/Medium/Low; evidence explains contributing signals."),
        ("CampaignRecord", "campaign_id, accounts, number_of_accounts, number_of_connections, average_coordination_score, review_status", "Calculates connected components via NetworkX; tags status as 'Requires Human Review' to prevent overstated certainty.")
    ]

    for i, row in enumerate(rows_data):
        for j, val in enumerate(row):
            cell = contract_table.rows[i+1].cells[j]
            set_cell_background(cell, "F8FAFC" if i % 2 == 0 else "FFFFFF")
            set_cell_margins(cell, 60, 60, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9.5)
            r.font.color.rgb = DARK_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # -------------------------------------------------------------
    # 3. Clean Ingestion Boundary (Stage 3)
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1_run = h1.add_run("3. Clean Ingestion Boundary (Stage 3)")
    h1_run.font.color.rgb = PRIMARY

    doc.add_paragraph(
        "A modular provider package was created in src/ingestion/ supporting three distinct ingestion modalities:"
    )

    doc.add_paragraph(
        "1. SyntheticDataProvider: Supplies the benchmark 10-post demonstration fixture. Every record is explicitly tagged provenance='synthetic_demo' to guarantee synthetic data is never confused with live intelligence.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "2. FileDataProvider: Imports authorized local CSV and JSON/JSONL posts and account metadata. Automatically infers missing account profiles directly from post activity.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "3. GenericAPIProvider: A production-ready REST adapter for authorized platform APIs and research feeds. Features server-side token masking (tokens are never logged), HTTP 429 rate-limit handling with Retry-After inspection, bounded exponential backoff retries, and pagination.",
        style='List Bullet'
    )

    # -------------------------------------------------------------
    # 4. Resilient NLP & Pipeline Unification (Stages 4 & 5)
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1_run = h1.add_run("4. Resilient NLP & End-to-End Pipeline Unification (Stages 4 & 5)")
    h1_run.font.color.rgb = PRIMARY

    doc.add_paragraph(
        "Two major architectural improvements were delivered in the core analysis engine:"
    )

    doc.add_paragraph(
        "• Dual-Engine Semantic Similarity (src/semantic_similarity.py): Implemented a resilient fallback mechanism. "
        "If sentence-transformers is installed, the system computes dense embeddings with all-MiniLM-L6-v2. "
        "If PyTorch is absent or the environment is offline, it automatically falls back to scikit-learn's TfidfVectorizer "
        "(word + character n-grams) and computes cosine similarity, logging a clear notice and maintaining full pipeline continuity.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "• Automated ML Bot Classification (src/account_analysis.py): Unified Phase 1 and Phase 2. "
        "Account behavioral metrics (total posts, URL ratio, duplicate text ratio, time between posts) are formatted into the "
        "21-feature schema and passed through models/selected_bot_classifier.joblib. The model outputs bot_probability and "
        "predicted_label_name for all analyzed accounts without altering the model file.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "• Interpretable Risk Scoring (src/risk_scoring.py): Coordination composite scores (based on text similarity, time proximity, and shared URLs) "
        "are strictly separated from bot classification probabilities. Evidence strings detail every observable signal so investigators can audit the score.",
        style='List Bullet'
    )

    # -------------------------------------------------------------
    # 5. Responsible Model Validation & Audit (Stage 6)
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1_run = h1.add_run("5. Responsible Model Audit & Generalization Report (Stage 6)")
    h1_run.font.color.rgb = PRIMARY

    doc.add_paragraph(
        "Per safety guidelines, the existing model artifact models/selected_bot_classifier.joblib was preserved intact. "
        "A dedicated audit script (src/ml/model_report.py) was built to inspect the artifact and generate docs/MODEL_EVALUATION_REPORT.md:"
    )

    metric_table = doc.add_table(rows=6, cols=3)
    metric_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    m_headers = ["Metric", "Test Set Value (N = 2,448)", "Evaluation Context & Protocol"]
    for j, h in enumerate(m_headers):
        cell = metric_table.rows[0].cells[j]
        set_cell_background(cell, "7C3AED")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(10)

    m_rows = [
        ("Accuracy", "0.9853 (98.5%)", "Stratified random split (70% train, 15% validation, 15% test)."),
        ("Precision (Bot)", "0.9963 (99.6%)", "Very low false positive rate on benchmark test accounts."),
        ("Recall (Bot)", "0.9816 (98.2%)", "Identifies 1,604 out of 1,634 benchmark bot accounts."),
        ("F1 Score (Bot)", "0.9889", "Selected over Logistic Regression (0.907) and Decision Tree (0.976)."),
        ("ROC-AUC", "0.9985", "Area under the Receiver Operating Characteristic curve.")
    ]

    for i, row in enumerate(m_rows):
        for j, val in enumerate(row):
            cell = metric_table.rows[i+1].cells[j]
            set_cell_background(cell, "F8FAFC" if i % 2 == 0 else "FFFFFF")
            set_cell_margins(cell, 60, 60, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9.5)
            r.font.color.rgb = DARK_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    p_gen = doc.add_paragraph()
    r_gen_bold = p_gen.add_run("Generalization Finding: ")
    r_gen_bold.font.bold = True
    p_gen.add_run(
        "Cross-group evaluation (check_group_generalization.py) shows that when testing on entirely unseen bot families (source groups), "
        "model recall drops significantly compared to random stratified splits. The classifier detects resemblance to historical training clusters, "
        "not a universal signature of automation. Predictions must never be presented as proof of malicious intent."
    )

    # -------------------------------------------------------------
    # 6. Streamlit Intelligence Dashboard Integration (Stage 7)
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1_run = h1.add_run("6. Streamlit Intelligence Dashboard Integration (Stage 7)")
    h1_run.font.color.rgb = PRIMARY

    doc.add_paragraph(
        "The 5-tab Streamlit dashboard (dashboard/app.py) was enhanced to connect directly to the unified pipeline results:"
    )

    doc.add_paragraph(
        "• Tab 1 (Posts): Shows ingested posts, keyword frequency charts, and search filters with clear demo data tagging.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "• Tab 2 (Coordination Detection): Displays flagged pairs, time difference vs. similarity scatter plot, and score distribution.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "• Tab 3 (Risk Scores): Renders investigation priority breakdown (High, Medium, Low) and transparent signal evidence.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "• Tab 4 (Campaign Analysis): Interactive Plotly network graph rendering candidate groups and account connections.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "• Tab 5 (Bot Classifier): Features the active dataset ML classification table alongside the model comparison and single-account prediction form.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "• Sidebar Utilities: Displays active data source, analysis run timestamp (UTC), custom CSV/JSON file uploader with one-click pipeline run, and live API connection status.",
        style='List Bullet'
    )

    # -------------------------------------------------------------
    # 7. Automated Test Verification & Release Checks (Stage 8)
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1_run = h1.add_run("7. Automated Test Verification & Release Checks (Stage 8)")
    h1_run.font.color.rgb = PRIMARY

    doc.add_paragraph(
        "The automated test suite was expanded significantly to cover all new contracts, providers, and pipeline stages:"
    )

    test_table = doc.add_table(rows=5, cols=3)
    test_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_headers = ["Test Suite File", "Test Count", "Functional Scope & Verification"]
    for j, h in enumerate(t_headers):
        cell = test_table.rows[0].cells[j]
        set_cell_background(cell, "0E7490")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(10)

    t_rows = [
        ("tests/test_backup.py", "12 Passed", "Verifies backup helper immutability, directory creation, byte-for-byte fidelity, and model exclusion."),
        ("tests/test_ingestion_and_contracts.py", "15 Passed", "Validates schema normalization, malformed post filtering, CSV/JSON loaders, mock API retries, rate-limits, token masking, and dry-run pipeline."),
        ("tests/test_ml_pipeline.py", "11 Passed", "Verifies label mapping, dataset discovery, feature extraction, train/test splitting, and model save/load."),
        ("tests/test_nextrace.py", "4 Passed", "Validates duplicate removal, invalid timestamp filtering, coordination detection, and risk scoring evidence.")
    ]

    for i, row in enumerate(t_rows):
        for j, val in enumerate(row):
            cell = test_table.rows[i+1].cells[j]
            set_cell_background(cell, "F8FAFC" if i % 2 == 0 else "FFFFFF")
            set_cell_margins(cell, 60, 60, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9.5)
            r.font.color.rgb = DARK_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    p_tot = doc.add_paragraph()
    r_tot = p_tot.add_run("Total Test Suite: 42 passed in 4.39 seconds (100% pass rate, 0 failures, 0 regressions).")
    r_tot.font.bold = True
    r_tot.font.color.rgb = PRIMARY

    # -------------------------------------------------------------
    # 8. Git & Safety Compliance Status
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1_run = h1.add_run("8. Git Status & Safety Compliance")
    h1_run.font.color.rgb = PRIMARY

    doc.add_paragraph(
        "• Branch: main (clean working tree prior to execution; all modifications are local and uncommitted per Rule A.3).",
        style='List Bullet'
    )
    doc.add_paragraph(
        "• Model Preservation: models/selected_bot_classifier.joblib was never retrained, renamed, or modified.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "• No Secrets Committed: Live credentials are kept out of Git. Sanitized template provided in .env.example.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "• Data Governance: No private datasets were uploaded or redistributed. Raw Cresci archives remain excluded.",
        style='List Bullet'
    )

    # Save Document
    docx_path = ROOT / "docs" / "NexTrace_Engineering_Continuation_Report.docx"
    doc.save(docx_path)
    print(f"Report successfully saved to: {docx_path}")
    return docx_path

if __name__ == "__main__":
    create_report()
